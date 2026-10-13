import io

from docx import Document
from docx.shared import Pt


def _is_heading(line: str) -> bool:
    letters = [c for c in line if c.isalpha()]
    return len(letters) >= 3 and all(c.isupper() for c in letters) and len(line) < 40


def resume_to_docx(text: str) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if _is_heading(line):
            doc.add_heading(line.title(), level=2)
        elif line[0] in "-•*":
            doc.add_paragraph(line.lstrip("-•* ").strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
