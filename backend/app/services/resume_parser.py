import io
import os

import pdfplumber
from docx import Document


class ResumeParseError(Exception):
    """Raised when we can't pull any usable text out of the uploaded file."""


def _extract_pdf(data: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    # Paragraphs cover most resumes; tables catch the ones laid out in a grid.
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    """Return the plain text of a resume uploaded as PDF or DOCX."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    else:
        raise ResumeParseError(f"Unsupported file type: {ext}")

    text = text.strip()
    # A scanned/image-only PDF parses to almost nothing — tell the user instead
    # of sending an empty resume to the LLM.
    if len(text) < 30:
        raise ResumeParseError(
            "Couldn't read any text from this file. If it's a scanned PDF, "
            "please upload a text-based resume."
        )
    return text
